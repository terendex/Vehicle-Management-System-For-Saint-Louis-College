"""The Excel export's layout and the report name casing.

The Excel file is a plain sortable table (no letterhead), but it should still
read like the PDF: nothing cut off at the next cell, long text wrapped rather
than stretching a column across the screen, and people's names in the same
"Juan Dela Cruz" form in both formats.
"""
import io

from django.test import SimpleTestCase
from openpyxl import load_workbook

from report_utils import branded_excel_response, name_case


class NameCaseTests(SimpleTestCase):
    def test_all_caps_names_are_title_cased(self):
        self.assertEqual(name_case('ALADIN C. VILLAREAL'), 'Aladin C. Villareal')
        self.assertEqual(name_case("D'ANGELO O'NEIL-SMITH JR."), "D'Angelo O'Neil-Smith Jr.")

    def test_generational_suffix_stays_upper(self):
        self.assertEqual(name_case('JUAN DELA CRUZ III'), 'Juan Dela Cruz III')

    def test_a_name_typed_in_mixed_case_is_left_alone(self):
        self.assertEqual(name_case('Ronald McDonald'), 'Ronald McDonald')

    def test_blank_stays_blank_for_the_callers_fallback(self):
        self.assertEqual(name_case(None), '')
        self.assertEqual(name_case('  '), '')


class ExcelLayoutTests(SimpleTestCase):
    HEADERS = ['#', 'Entry Gate', 'Remarks']

    def sheet(self, rows, widths=(6, 10, 30)):
        resp = branded_excel_response(filename='x.xlsx', sheet_title='Log', report_title='t',
                                      subtitle='s', headers=self.HEADERS, rows=rows,
                                      col_widths=list(widths))
        return load_workbook(io.BytesIO(resp.content)).active

    def test_a_column_widens_to_fit_its_longest_value(self):
        ws = self.sheet([[1, 'Gate 1 — Main Entrance', 'ok']])
        self.assertGreaterEqual(ws.column_dimensions['B'].width, len('Gate 1 — Main Entrance'))

    def test_the_callers_width_is_a_floor(self):
        ws = self.sheet([[1, 'G1', 'ok']], widths=(6, 18, 30))
        self.assertEqual(ws.column_dimensions['B'].width, 18)

    def test_long_text_wraps_at_the_cap_and_short_text_does_not(self):
        ws = self.sheet([[1, 'G1', 'x' * 200], [2, 'G1', 'short']])
        self.assertEqual(ws.column_dimensions['C'].width, 60)
        self.assertTrue(ws['C2'].alignment.wrap_text)
        self.assertFalse(ws['C3'].alignment.wrap_text)

    def test_header_row_is_frozen_and_filterable(self):
        ws = self.sheet([[1, 'G1', 'a'], [2, 'G1', 'b']])
        self.assertEqual(ws.freeze_panes, 'A2')
        self.assertEqual(ws.auto_filter.ref, 'A1:C3')

    def test_an_empty_report_still_has_a_filter_range(self):
        ws = self.sheet([])
        self.assertEqual(ws['A1'].value, '#')
        self.assertEqual(ws.auto_filter.ref, 'A1:C2')
